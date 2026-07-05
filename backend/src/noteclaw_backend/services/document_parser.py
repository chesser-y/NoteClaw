from __future__ import annotations

import csv
from dataclasses import dataclass, field
from functools import lru_cache
from io import BytesIO, StringIO
import os
from pathlib import Path
import shutil
import sys

from PIL import Image, ImageOps
import pytesseract

from noteclaw_backend.domain.enums import ContentType


@dataclass(frozen=True)
class ParsedDocument:
    content_type: ContentType
    content: str
    title: str | None = None
    metadata: dict = field(default_factory=dict)


CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".go", ".rs", ".java", ".cpp", ".c",
    ".h", ".cs", ".php", ".rb", ".swift", ".kt", ".sql", ".sh", ".json", ".yaml", ".yml",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
TABLE_EXTENSIONS = {".csv", ".tsv"}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".html", ".htm"}


def parse_file_bytes(
    *,
    filename: str,
    data: bytes,
    content_type: ContentType | None = None,
    saved_path: Path | None = None,
) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    inferred = content_type or infer_content_type(filename)
    title = Path(filename).stem or filename
    metadata = {"filename": filename, "extension": suffix}
    if saved_path is not None:
        metadata["stored_path"] = str(saved_path)

    if inferred == ContentType.IMAGE or suffix in IMAGE_EXTENSIONS:
        ocr_text = _ocr_image(data)
        content = "Image OCR text:\n" + (ocr_text.strip() or "No OCR text extracted.")
        return ParsedDocument(ContentType.IMAGE, content, title=title, metadata={**metadata, "ocr_text": ocr_text})

    if inferred == ContentType.TABLE or suffix in TABLE_EXTENSIONS:
        text = decode_bytes(data)
        table = delimited_text_to_markdown(text, delimiter="\t" if suffix == ".tsv" else None)
        return ParsedDocument(ContentType.TABLE, table or text, title=title, metadata=metadata)

    if suffix == ".pptx":
        ppt_text = _extract_pptx_text(data)
        if ppt_text:
            return ParsedDocument(ContentType.DOCUMENT, ppt_text, title=title, metadata=metadata)

    decoded = decode_bytes(data)
    if inferred == ContentType.CODE:
        return ParsedDocument(ContentType.CODE, decoded, title=title, metadata=metadata)
    if suffix in TEXT_EXTENSIONS:
        return ParsedDocument(ContentType.TEXT, decoded, title=title, metadata=metadata)
    return ParsedDocument(inferred, decoded, title=title, metadata=metadata)


def infer_content_type(filename: str) -> ContentType:
    suffix = Path(filename).suffix.lower()
    if suffix in CODE_EXTENSIONS:
        return ContentType.CODE
    if suffix in IMAGE_EXTENSIONS:
        return ContentType.IMAGE
    if suffix in TABLE_EXTENSIONS:
        return ContentType.TABLE
    if suffix in TEXT_EXTENSIONS:
        return ContentType.TEXT
    return ContentType.DOCUMENT


def decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def delimited_text_to_markdown(text: str, delimiter: str | None = None) -> str:
    sample = text[:2048]
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(sample)
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _ocr_image(data: bytes) -> str:
    try:
        _configure_tesseract()
        image = _prepare_ocr_image(Image.open(BytesIO(data)))
        text = _run_tesseract(image, config="--psm 6")
        if not text.strip():
            text = _run_tesseract(image, config="--psm 11")
        return text
    except Exception:
        return ""


@lru_cache(maxsize=1)
def _configure_tesseract() -> None:
    candidates = [
        os.environ.get("TESSERACT_CMD"),
        str(Path(sys.prefix) / "bin" / "tesseract"),
        shutil.which("tesseract"),
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            pytesseract.pytesseract.tesseract_cmd = candidate
            return


def _prepare_ocr_image(image: Image.Image) -> Image.Image:
    image = ImageOps.exif_transpose(image)
    if image.mode not in {"RGB", "L"}:
        image = image.convert("RGB")

    max_side = max(image.size)
    if max_side > 1800:
        scale = 1800 / max_side
        size = (max(1, int(image.width * scale)), max(1, int(image.height * scale)))
        image = image.resize(size, Image.Resampling.LANCZOS)
    elif max_side < 1000:
        image = image.resize((image.width * 2, image.height * 2), Image.Resampling.LANCZOS)

    gray = ImageOps.grayscale(image)
    return ImageOps.autocontrast(gray)


def _run_tesseract(image: Image.Image, *, config: str) -> str:
    languages = _ocr_languages()
    errors: list[Exception] = []
    for language in languages:
        try:
            return pytesseract.image_to_string(image, lang=language, config=config)
        except Exception as exc:
            errors.append(exc)
    if errors:
        raise errors[-1]
    return ""


@lru_cache(maxsize=1)
def _ocr_languages() -> tuple[str | None, ...]:
    try:
        available = set(pytesseract.get_languages(config=""))
    except Exception:
        return (None,)

    preferred: list[str | None] = []
    if {"chi_sim", "eng"}.issubset(available):
        preferred.append("chi_sim+eng")
    if "eng" in available:
        preferred.append("eng")
    if "chi_sim" in available:
        preferred.append("chi_sim")
    return tuple(preferred) or (None,)


def _extract_pptx_text(data: bytes) -> str:
    try:
        from pptx import Presentation
    except Exception:
        return ""
    try:
        presentation = Presentation(BytesIO(data))
    except Exception:
        return ""
    slides: list[str] = []
    for idx, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                parts.append(text.strip())
        if parts:
            slides.append(f"Slide {idx}\n" + "\n".join(parts))
    return "\n\n".join(slides)
